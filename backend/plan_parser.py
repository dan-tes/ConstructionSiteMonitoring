"""Block 1 — turn the uploaded plan workbook into `plan_stages` for phase/delay.

Parses the canonical plan format: see
phase_determination/data/Требования_к_каноническому_формату_плана_v2.docx.
The canonical table is activity-level (one row per construction activity,
the same shape as phase_determination/data/activities_with_equipment.csv)
with a fixed 21-column header — `project_id`, `activity_id`, `phase`,
`phase_order`, `activity_sequence`, `activity_name`, `resource_type`,
`planned_duration_days`, `quantity`, `unit_cost`, `planned_cost`,
`criticality`, `status`, `critical_path`, `critical_path_position`,
`project_network_duration`, `predecessor_count`, `successor_count`, `split`,
`expected_equipment`, `expected_equipment_descriptions`. Notably, no
`planned_start`/`planned_end` — the canonical format is durations-only (see
the docx §22).

This module extracts two things:

- `phases`: `phase`/`phase_order`/`planned_duration_days` (summed across
  that phase's activities). Activities inside a phase commonly run in
  parallel, so this sum is *not* the phase's real elapsed calendar time —
  it's the same aggregation the phase-detection training notebook uses for
  its (purely semantic) phase-embedding features, reused here because
  nothing more precise exists in the canonical format. Fine as a relative
  weighting between phases; not fine as an absolute day count on its own.
- `project_duration_days`: `project_network_duration`, the actual
  critical-path project timeline (one value, repeated on every row of the
  project — see docx §13, which calls this field out as kept specifically
  "for subsequent analysis and delay forecasting"). This is what
  `delay/worker.py` rescales the phase sums against, so offsets end up on
  the real project timeline instead of the inflated activity-sum one — see
  its module docstring for the rescaling.

Delivered as one `.xlsx` workbook, the activities table on one sheet among
others (equipment reference sheets, etc.) — this looks for a sheet named
like "plan"/"activities"/"activities_with_equipment" and falls back to the
first sheet if none match, since only the activities table matters here.

Deliberately minimal, same spirit as the rest of block 1: the workbook is
assumed to already be in the canonical shape (see integrations/README.md —
"assume we're handed data that's ideal for parsing"). No column-name
fuzzing, no unit conversion, no `.csv`/`.xls` support (the canonical format
is delivered as a multi-sheet `.xlsx`) — an out-of-shape file fails loudly
via `PlanParseError` instead of guessing.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from integrations.schemas import PlanPhaseIn

CANONICAL_COLUMNS = [
    "project_id",
    "activity_id",
    "phase",
    "phase_order",
    "activity_sequence",
    "activity_name",
    "resource_type",
    "planned_duration_days",
    "quantity",
    "unit_cost",
    "planned_cost",
    "criticality",
    "status",
    "critical_path",
    "critical_path_position",
    "project_network_duration",
    "predecessor_count",
    "successor_count",
    "split",
    "expected_equipment",
    "expected_equipment_descriptions",
]

_PLAN_SHEET_NAMES = {"plan", "activities", "activities_with_equipment", "activity"}


class PlanParseError(ValueError):
    """The plan workbook isn't in the expected canonical shape."""


@dataclass
class ParsedPlan:
    phases: list[PlanPhaseIn]
    project_duration_days: float


def _find_plan_sheet(workbook: openpyxl.Workbook):
    for name in workbook.sheetnames:
        key = name.strip().lower().replace(" ", "_")
        if key in _PLAN_SHEET_NAMES or key.startswith("activities"):
            return workbook[name]
    return workbook[workbook.sheetnames[0]]


def parse_plan_workbook(path: Path) -> ParsedPlan:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise PlanParseError(f"cannot open plan workbook {path}: {exc}") from exc

    sheet = _find_plan_sheet(workbook)
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise PlanParseError(f"sheet {sheet.title!r} is empty")

    header = [str(h).strip() if h is not None else "" for h in header]
    missing = [c for c in CANONICAL_COLUMNS if c not in header]
    if missing:
        raise PlanParseError(
            f"sheet {sheet.title!r} is missing canonical columns {missing} "
            f"(got {header})"
        )
    col = {name: header.index(name) for name in CANONICAL_COLUMNS}

    project_ids: set[str] = set()
    phase_durations: dict[tuple[int, str], float] = {}
    network_durations: set[float] = set()

    for row_num, row in enumerate(rows, start=2):  # header is row 1
        if row is None or all(v is None for v in row):
            continue
        project_id = row[col["project_id"]]
        phase = row[col["phase"]]
        phase_order = row[col["phase_order"]]
        duration = row[col["planned_duration_days"]]
        network_duration = row[col["project_network_duration"]]
        if (
            project_id is None
            or phase is None
            or phase_order is None
            or duration is None
            or network_duration is None
        ):
            raise PlanParseError(
                f"row {row_num}: missing project_id/phase/phase_order/"
                "planned_duration_days/project_network_duration"
            )
        project_ids.add(str(project_id))
        network_durations.add(float(network_duration))
        key = (int(phase_order), str(phase).strip())
        phase_durations[key] = phase_durations.get(key, 0.0) + float(duration)

    if not phase_durations:
        raise PlanParseError(f"sheet {sheet.title!r} has no data rows")
    if len(project_ids) > 1:
        raise PlanParseError(
            f"sheet {sheet.title!r} has multiple project_id values {sorted(project_ids)} — "
            "expected exactly one project per plan upload"
        )
    if len(network_durations) > 1:
        raise PlanParseError(
            f"sheet {sheet.title!r} has inconsistent project_network_duration values "
            f"{sorted(network_durations)} for the same project"
        )

    phases = [
        PlanPhaseIn(phase=phase, phase_order=order, planned_duration_days=duration)
        for (order, phase), duration in sorted(phase_durations.items())
    ]
    return ParsedPlan(phases=phases, project_duration_days=next(iter(network_durations)))


@dataclass
class CanonicalExportPhase:
    phase: str
    phase_order: int
    planned_duration_days: float
    expected_equipment: list[str]


def build_canonical_workbook(
    project_id: str,
    phases: list[CanonicalExportPhase],
    project_duration_days: float,
    statuses: dict[str, str] | None = None,
) -> openpyxl.Workbook:
    """The write side of the canonical format: one `.xlsx` with the same
    21-column header `parse_plan_workbook` reads, so re-uploading this file
    round-trips cleanly through it. One synthetic activity row stands in for
    each whole phase — `PlanStage` (see models.py) only ever has phase-level
    durations, not the real per-activity breakdown a genuine canonical plan
    has, whether the project's plan came from the fast path or the
    `planner` service's LLM fallback. Lets a project be inspected/edited in
    the canonical shape and re-uploaded, rather than only ever consumed
    read-only.

    `statuses` fills the `status` column per phase (see progress.py — the
    live "fact" export); phases missing from it stay "Planned". Returns the
    open workbook rather than bytes so progress.py can append its own
    sheets — `parse_plan_workbook` only ever reads the activities sheet.
    """
    statuses = statuses or {}
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "activities"
    sheet.append(CANONICAL_COLUMNS)
    last = len(phases) - 1
    for i, p in enumerate(phases):
        sheet.append(
            [
                project_id,
                f"{project_id}_P{p.phase_order:02d}",
                p.phase,
                p.phase_order,
                1,  # activity_sequence — one synthetic activity per phase
                p.phase,  # activity_name
                "mixed",  # resource_type
                p.planned_duration_days,
                1,  # quantity
                0,  # unit_cost
                0,  # planned_cost
                1.0,  # criticality — this row *is* the whole phase
                statuses.get(p.phase, "Planned"),  # status
                1,  # critical_path
                p.phase_order,  # critical_path_position
                project_duration_days,
                0 if i == 0 else 1,  # predecessor_count
                0 if i == last else 1,  # successor_count
                "validation",  # split
                repr(p.expected_equipment),
                "{}",  # expected_equipment_descriptions — not tracked per-phase
            ]
        )
    return workbook


def workbook_bytes(workbook: openpyxl.Workbook) -> bytes:
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()
